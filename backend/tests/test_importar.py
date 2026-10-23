from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from models.categoria_model import Categoria
from models.evento_model import Evento
from models.produto_model import Produto
from services.importar_planilha_service import MODULOS
from tests.fabricas import criar_categoria, criar_produto


def enviar(client, modulo, conteudo, nome="planilha.csv"):
    if isinstance(conteudo, str):
        conteudo = conteudo.encode("utf-8")
    return client.post(
        f"/importar?modulo={modulo}",
        data={"arquivo": (BytesIO(conteudo), nome)},
        content_type="multipart/form-data",
    )


@pytest.mark.parametrize("modulo", list(MODULOS))
def test_baixa_modelo_xlsx(client, modulo):
    r = client.get(f"/importar/modelo?modulo={modulo}")
    assert r.status_code == 200
    assert r.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert f"modelo_{modulo}.xlsx" in r.headers["Content-Disposition"]
    assert r.data[:2] == b"PK"
    planilha = load_workbook(BytesIO(r.data)).active
    cabecalho = [c.value for c in planilha[1]]
    assert cabecalho == MODULOS[modulo]["colunas"]


def test_modelo_de_modulo_invalido(client):
    r = client.get("/importar/modelo?modulo=clientes")
    assert r.status_code == 400
    assert r.get_json() == {"erro": "módulo inválido. use: categorias, produtos, vendas ou eventos"}


def test_importa_produtos_csv_com_ponto_e_virgula(client):
    cat = criar_categoria("Laticínios")
    csv = (
        "nome;id_categoria;codigo_barras;preco_custo;preco_venda;quantidade_estoque;estoque_minimo;unidade;data_validade\n"
        f"Queijo Minas;{cat.id_categoria};789001;12,00;19,90;20;5;un;10/09/2026\n"
        "Sem preço;;789002;;;1;;;\n"
    )
    r = enviar(client, "produtos", csv)
    assert r.status_code == 200
    assert r.get_json() == {"importados": 1, "ignorados": 0, "erros": ["Linha 3: preco_venda é obrigatório"]}

    queijo = Produto.buscar_por_codigo("789001")
    assert queijo.nome == "Queijo Minas"
    assert float(queijo.preco_venda) == 19.9
    assert queijo.unidade == "UN"
    assert queijo.data_validade.isoformat() == "2026-09-10"


def test_importa_produto_existente_atualiza_pelo_codigo(client):
    criar_produto("Nome antigo", codigo_barras="789001", quantidade_estoque=1)
    csv = "nome,codigo_barras,preco_venda,quantidade_estoque\nNome novo,789001,9.90,15\n"
    assert enviar(client, "produtos", csv).get_json()["importados"] == 1
    p = Produto.buscar_por_codigo("789001")
    assert p.nome == "Nome novo"
    assert p.quantidade_estoque == 15
    assert Produto.query.count() == 1


def test_importa_produtos_xlsx(client):
    wb = Workbook()
    ws = wb.active
    ws.append(["nome", "preco_venda", "quantidade_estoque"])
    ws.append(["Manteiga", 8.5, 12])
    arquivo = BytesIO()
    wb.save(arquivo)

    r = enviar(client, "produtos", arquivo.getvalue(), nome="produtos.xlsx")
    assert r.get_json() == {"importados": 1, "ignorados": 0, "erros": []}
    assert Produto.query.filter_by(nome="Manteiga").first().quantidade_estoque == 12


def test_importa_vendas_baixa_estoque(client):
    p = criar_produto("Leite", quantidade_estoque=10)
    csv = f"id_produto,quantidade,data_venda\n{p.id_produto},3,2026-01-05\n999,1,2026-01-05\n"
    r = enviar(client, "vendas", csv)
    assert r.get_json() == {"importados": 1, "ignorados": 0, "erros": ["Linha 3: produto não existe"]}
    assert Produto.buscar_por_id(p.id_produto).quantidade_estoque == 7


def test_importa_categorias_ignora_repetida_e_linha_vazia(client):
    criar_categoria("Bebidas")
    csv = "nome,descricao\nBebidas,repetida\nPadaria,pães\n\n"
    r = enviar(client, "categorias", csv)
    # ignorados: a categoria repetida e a linha em branco
    assert r.get_json() == {"importados": 1, "ignorados": 2, "erros": []}
    assert {c.nome for c in Categoria.query.all()} == {"Bebidas", "Padaria"}


def test_importa_eventos(client):
    csv = "nome,data_inicio,data_fim,impacto_percentual\nFeira,2026-09-01,2026-09-03,20\nRuim,2026-09-05,2026-09-01,0\n"
    r = enviar(client, "eventos", csv)
    assert r.get_json() == {
        "importados": 1,
        "ignorados": 0,
        "erros": ["Linha 3: data_fim tem que ser depois da data_inicio"],
    }
    assert Evento.query.one().impacto_percentual == 20


def test_importa_sem_arquivo(client):
    r = client.post("/importar?modulo=produtos", data={}, content_type="multipart/form-data")
    assert r.status_code == 400
    assert r.get_json() == {"erro": "envie um arquivo .csv ou .xlsx"}


@pytest.mark.parametrize("nome, conteudo, erro", [
    ("produtos.txt", "nome,preco_venda\nX,1\n", "o arquivo tem que ser .csv ou .xlsx"),
    ("produtos.csv", "", "arquivo vazio"),
    ("produtos.csv", "nome,quantidade_estoque\nX,1\n", "faltam as colunas obrigatórias: preco_venda"),
])
def test_importa_arquivo_invalido(client, nome, conteudo, erro):
    r = enviar(client, "produtos", conteudo, nome=nome)
    assert r.status_code == 400
    assert r.get_json() == {"erro": erro}


def test_importa_modulo_invalido(client):
    r = enviar(client, "clientes", "nome\nX\n")
    assert r.status_code == 400
