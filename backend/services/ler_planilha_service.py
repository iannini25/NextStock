import csv
from io import BytesIO, StringIO
from openpyxl import load_workbook
from services.formato_planilha import _texto, _cabecalho


class LerPlanilhaService:

    def executar(self, nome, dados):
        return self._ler_csv(dados) if nome.endswith(".csv") else self._ler_xlsx(dados)

    def _ler_csv(self, dados):
        try:
            texto = dados.decode("utf-8-sig")
        except UnicodeDecodeError:
            texto = dados.decode("cp1252")
        try:
            dialect = csv.Sniffer().sniff(texto[:2048], delimiters=",;")
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.reader(StringIO(texto), dialect))
        return self._linhas_de_tabela(rows)

    def _ler_xlsx(self, dados):
        try:
            wb = load_workbook(BytesIO(dados), data_only=True, read_only=True)
        except Exception:
            raise ValueError("não foi possível ler o arquivo xlsx")
        ws = wb.active
        rows = [list(row) for row in ws.iter_rows(values_only=True)]
        wb.close()
        return self._linhas_de_tabela(rows)

    def _linhas_de_tabela(self, rows):
        if not rows:
            raise ValueError("arquivo sem cabeçalho")
        cabecalhos = [_cabecalho(c) for c in rows[0]]
        if not any(cabecalhos):
            raise ValueError("arquivo sem cabeçalho")
        linhas = []
        for row in rows[1:]:
            item = {}
            for i, chave in enumerate(cabecalhos):
                if not chave:
                    continue
                item[chave] = _texto(row[i] if i < len(row) else None)
            linhas.append(item)
        return linhas, {c for c in cabecalhos if c}
