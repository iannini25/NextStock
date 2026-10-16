import csv
from datetime import date, datetime
from io import BytesIO, StringIO

from openpyxl import Workbook, load_workbook

from models.database import db
from models.categoria_model import Categoria
from models.evento_model import Evento
from models.produto_model import Produto
from services.estoque_lote_service import EstoqueLoteService
from services.registrar_venda_service import RegistrarVendaService
from services.utils import inteiro, numero, virar_data

# so as colunas que o sistema aceita em cada modulo
MODULOS = {
    "categorias": {
        "colunas": ["nome", "descricao"],
        "obrigatorias": ["nome"],
        "exemplo": ["Bebidas", "Refrigerantes e sucos"],
    },
    "produtos": {
        "colunas": [
            "nome",
            "id_categoria",
            "codigo_barras",
            "preco_custo",
            "preco_venda",
            "quantidade_estoque",
            "estoque_minimo",
            "unidade",
            "data_validade",
        ],
        "obrigatorias": ["nome", "preco_venda"],
        "exemplo": [
            "Queijo Minas 500g",
            "1",
            "7890000000062",
            "12.00",
            "19.90",
            "20",
            "5",
            "UN",
            "2026-09-10",
        ],
    },
    "vendas": {
        "colunas": ["id_produto", "quantidade", "data_venda"],
        "obrigatorias": ["id_produto", "quantidade"],
        "exemplo": ["1", "3", "2026-08-20"],
    },
    "eventos": {
        "colunas": ["nome", "data_inicio", "data_fim", "impacto_percentual"],
        "obrigatorias": ["nome", "data_inicio", "data_fim"],
        "exemplo": ["Festival local", "2026-09-01", "2026-09-03", "20"],
    },
}


def _texto(v):
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v).strip()


def _cabecalho(v):
    return _texto(v).lower()


def _numero(v):
    t = _texto(v).replace(" ", "").replace(",", ".")
    return numero(t) if t else None


def _data(v):
    t = _texto(v)
    if not t:
        return None
    if "T" in t:
        t = t.split("T")[0]
    if " " in t:
        t = t.split(" ")[0]
    if "/" in t:
        partes = t.split("/")
        if len(partes) == 3:
            d, m, a = partes
            if len(a) == 4:
                t = f"{a}-{m.zfill(2)}-{d.zfill(2)}"
    return virar_data(t)


class ImportarPlanilhaService:

    def gerar_modelo(self, modulo):
        modulo = (modulo or "").strip().lower()
        cfg = self._config(modulo)
        wb = Workbook()
        ws = wb.active
        ws.title = modulo
        ws.append(cfg["colunas"])
        ws.append(cfg["exemplo"])
        bio = BytesIO()
        wb.save(bio)
        bio.seek(0)
        return bio, f"modelo_{modulo}.xlsx"

    def executar(self, modulo, arquivo):
        modulo = (modulo or "").strip().lower()
        cfg = self._config(modulo)
        if arquivo is None or not arquivo.filename:
            raise ValueError("envie um arquivo .csv ou .xlsx")

        nome = arquivo.filename.lower()
        if not (nome.endswith(".csv") or nome.endswith(".xlsx")):
            raise ValueError("o arquivo tem que ser .csv ou .xlsx")

        dados = arquivo.read()
        if not dados:
            raise ValueError("arquivo vazio")

        linhas, cabecalhos = self._ler_arquivo(nome, dados)
        faltando = [c for c in cfg["obrigatorias"] if c not in cabecalhos]
        if faltando:
            raise ValueError("faltam as colunas obrigatórias: " + ", ".join(faltando))

        importados = 0
        ignorados = 0
        erros = []

        for i, item in enumerate(linhas, start=2):
            if self._linha_vazia(item, cfg["colunas"]):
                ignorados += 1
                continue
            try:
                resultado = self._importar_linha(modulo, item)
                if resultado == "ok":
                    importados += 1
                else:
                    ignorados += 1
            except ValueError as e:
                erros.append("Linha " + str(i) + ": " + str(e))

        return {
            "importados": importados,
            "ignorados": ignorados,
            "erros": erros,
        }

    def _config(self, modulo):
        chave = (modulo or "").strip().lower()
        if chave not in MODULOS:
            raise ValueError("módulo inválido. use: categorias, produtos, vendas ou eventos")
        return MODULOS[chave]

    def _ler_arquivo(self, nome, dados):
        if nome.endswith(".csv"):
            return self._ler_csv(dados)
        return self._ler_xlsx(dados)

    def _ler_csv(self, dados):
        try:
            texto = dados.decode("utf-8-sig")
        except UnicodeDecodeError:
            texto = dados.decode("cp1252")
        try:
            dialect = csv.Sniffer().sniff(texto[:2048], delimiters=",;")
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.reader(StringIO(texto), dialect))
        return self._linhas_de_tabela(rows)

    def _ler_xlsx(self, dados):
        try:
            wb = load_workbook(BytesIO(dados), data_only=True, read_only=True)
        except Exception:
            raise ValueError("não foi possível ler o arquivo xlsx")
        ws = wb.active
        rows = [list(row) for row in ws.iter_rows(values_only=True)]
        wb.close()
        return self._linhas_de_tabela(rows)

    def _linhas_de_tabela(self, rows):
        if not rows:
            raise ValueError("arquivo sem cabeçalho")
        cabecalhos = [_cabecalho(c) for c in rows[0]]
        if not any(cabecalhos):
            raise ValueError("arquivo sem cabeçalho")
        linhas = []
        for row in rows[1:]:
            item = {}
            for i, chave in enumerate(cabecalhos):
                if not chave:
                    continue
                item[chave] = _texto(row[i] if i < len(row) else None)
            linhas.append(item)
        return linhas, {c for c in cabecalhos if c}

    def _linha_vazia(self, item, colunas):
        return all(not _texto(item.get(c)) for c in colunas)

    def _importar_linha(self, modulo, item):
        if modulo == "categorias":
            return self._importar_categoria(item)
        if modulo == "produtos":
            return self._importar_produto(item)
        if modulo == "vendas":
            return self._importar_venda(item)
        return self._importar_evento(item)

    def _importar_categoria(self, item):
        nome = _texto(item.get("nome"))
        if not nome:
            raise ValueError("nome é obrigatório")
        if Categoria.query.filter_by(nome=nome).first():
            return "skip"
        descricao = _texto(item.get("descricao")) or None
        Categoria(nome=nome, descricao=descricao).salvar()
        return "ok"

    def _importar_produto(self, item):
        nome = _texto(item.get("nome"))
        if not nome:
            raise ValueError("nome é obrigatório")
        preco_venda = _numero(item.get("preco_venda"))
        if preco_venda is None:
            raise ValueError("preco_venda é obrigatório")

        id_categoria = inteiro(_texto(item.get("id_categoria")))
        if id_categoria and not Categoria.buscar_por_id(id_categoria):
            raise ValueError("categoria não existe")

        codigo = _texto(item.get("codigo_barras")) or None
        produto = Produto.buscar_por_codigo(codigo) if codigo else None
        custo = _numero(item.get("preco_custo"))
        # planilha sem a coluna de quantidade: produto que ja existe fica com o estoque que tem
        tem_quantidade = "quantidade_estoque" in item
        qtd = inteiro(_texto(item.get("quantidade_estoque"))) or 0
        minimo = inteiro(_texto(item.get("estoque_minimo"))) or 0
        unidade = (_texto(item.get("unidade")) or "UN").upper()
        validade = _data(item.get("data_validade"))

        estoque = EstoqueLoteService()
        if not produto:
            produto = Produto(nome=nome, codigo_barras=codigo, quantidade_estoque=0)
            db.session.add(produto)
            tem_quantidade = qtd > 0
        produto.nome = nome
        produto.id_categoria = id_categoria
        produto.codigo_barras = codigo
        produto.preco_custo = custo
        produto.preco_venda = preco_venda
        produto.estoque_minimo = minimo
        produto.unidade = unidade
        try:
            # a quantidade da planilha vira ajuste nos lotes (sobrou = lote novo com a validade da planilha)
            if tem_quantidade:
                estoque.ajustar_para(produto, qtd, data_validade=validade, origem="importacao")
            db.session.commit()
        except Exception:
            db.session.rollback()
            raise
        return "ok"

    def _importar_venda(self, item):
        id_produto = inteiro(_texto(item.get("id_produto")))
        if not id_produto:
            raise ValueError("id_produto é obrigatório")
        if not Produto.buscar_por_id(id_produto):
            raise ValueError("produto não existe")
        quantidade = inteiro(_texto(item.get("quantidade")))
        if not quantidade:
            raise ValueError("quantidade tem que ser maior que zero")
        data_venda = _data(item.get("data_venda")) or date.today()
        RegistrarVendaService().executar({
            "id_produto": id_produto,
            "quantidade": quantidade,
            "data_venda": data_venda.isoformat(),
        })
        return "ok"

    def _importar_evento(self, item):
        nome = _texto(item.get("nome"))
        if not nome:
            raise ValueError("nome é obrigatório")
        data_inicio = _data(item.get("data_inicio"))
        data_fim = _data(item.get("data_fim"))
        if not data_inicio or not data_fim:
            raise ValueError("data_inicio e data_fim são obrigatórias")
        if data_fim < data_inicio:
            raise ValueError("data_fim tem que ser depois da data_inicio")
        try:
            impacto = int(_texto(item.get("impacto_percentual")) or 0)
        except ValueError:
            raise ValueError("impacto_percentual tem que ser número inteiro")
        Evento(
            nome=nome,
            data_inicio=data_inicio,
            data_fim=data_fim,
            impacto_percentual=impacto,
        ).salvar()
        return "ok"
